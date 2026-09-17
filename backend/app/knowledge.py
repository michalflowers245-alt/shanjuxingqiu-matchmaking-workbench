from __future__ import annotations

import csv
import hashlib
import io
import ipaddress
import json
import math
import mimetypes
import re
import socket
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

from .config import UPLOAD_DIR
from .db import Database, database, json_dumps, json_loads, new_id, utc_now
from .providers import ModelGateway, ProviderError, model_gateway


SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".html", ".htm", ".csv", ".tsv", ".pdf", ".docx", ".xlsx"}


def content_hash(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text).strip()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def safe_filename(name: str) -> str:
    clean = re.sub(r"[^\w\-. ()\u4e00-\u9fff]", "_", Path(name).name)
    return clean[:180] or "document.txt"


def chunk_text(text: str, size: int = 900, overlap: int = 120) -> list[str]:
    clean = re.sub(r"\r\n?", "\n", text)
    clean = re.sub(r"\n{3,}", "\n\n", clean).strip()
    if not clean:
        return []
    chunks: list[str] = []
    cursor = 0
    while cursor < len(clean):
        end = min(len(clean), cursor + size)
        if end < len(clean):
            boundary = max(clean.rfind("\n", cursor + size // 2, end), clean.rfind("。", cursor + size // 2, end))
            if boundary > cursor:
                end = boundary + 1
        piece = clean[cursor:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(clean):
            break
        cursor = max(cursor + 1, end - overlap)
    return chunks


def assert_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("只允许不含账号信息的公开 HTTP/HTTPS 地址")
    host = parsed.hostname.lower()
    if host in {"localhost", "localhost.localdomain"} or host.endswith(".local"):
        raise ValueError("不允许访问本机或局域网地址")
    try:
        addresses = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror as exc:
        raise ValueError("网址域名无法解析") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise ValueError("不允许访问本机、局域网或保留地址")


def fetch_public_page(url: str) -> tuple[str, str, str]:
    current = url
    headers = {"User-Agent": "Mozilla/5.0 CopyWorkbench/1.0"}
    with httpx.Client(timeout=25, follow_redirects=False, headers=headers) as client:
        for _ in range(4):
            assert_public_url(current)
            response = client.get(current)
            if response.is_redirect:
                location = response.headers.get("location")
                if not location:
                    raise ValueError("网页重定向缺少目标地址")
                current = urljoin(current, location)
                continue
            response.raise_for_status()
            content_type = response.headers.get("content-type", "text/html").split(";", 1)[0]
            if int(response.headers.get("content-length") or 0) > 10 * 1024 * 1024:
                raise ValueError("网页内容超过 10MB")
            if content_type == "application/pdf":
                text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(response.content)).pages)
                return current, Path(urlparse(current).path).name or "网页 PDF", text
            soup = BeautifulSoup(response.text, "html.parser")
            for node in soup(["script", "style", "noscript", "svg"]):
                node.decompose()
            title = soup.title.get_text(" ", strip=True) if soup.title else current
            text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
            return current, title[:200], text
    raise ValueError("网页重定向次数过多")


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    mime = mimetypes.guess_type(filename)[0] or "application/octet-stream"
    if suffix in {".txt", ".md", ".markdown"}:
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                return data.decode(encoding), mime
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace"), mime
    if suffix in {".html", ".htm"}:
        soup = BeautifulSoup(data, "html.parser")
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        return soup.get_text("\n", strip=True), mime
    if suffix in {".csv", ".tsv"}:
        raw = data.decode("utf-8-sig", errors="replace")
        delimiter = "\t" if suffix == ".tsv" else ","
        rows = csv.reader(io.StringIO(raw), delimiter=delimiter)
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows), mime
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages), mime
    if suffix == ".docx":
        document = DocxDocument(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return "\n".join(parts), mime
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows: list[str] = []
        for sheet in workbook.worksheets:
            rows.append(f"工作表：{sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    rows.append(" | ".join(values))
        return "\n".join(rows), mime
    raise ValueError(f"暂不支持文件类型：{suffix or '未知'}")


class KnowledgeService:
    def __init__(self, db: Database = database, models: ModelGateway = model_gateway):
        self.db = db
        self.models = models

    def default_library(self, workspace_id: str) -> str:
        row = self.db.one("SELECT id FROM libraries WHERE workspace_id=? ORDER BY created_at LIMIT 1", (workspace_id,))
        if not row:
            raise ValueError("品牌资料库不存在")
        return row["id"]

    def import_text(
        self,
        workspace_id: str,
        title: str,
        text: str,
        source_type: str = "paste",
        source_uri: str = "",
        library_id: str | None = None,
        mime_type: str = "text/plain",
        is_gold: bool = False,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        clean = text.strip()
        if not clean:
            raise ValueError("资料没有可读取的文字")
        digest = content_hash(clean)
        duplicate = self.db.one(
            "SELECT * FROM documents WHERE workspace_id=? AND content_hash=?",
            (workspace_id, digest),
        )
        if duplicate:
            return {**duplicate, "duplicate": True}
        library = library_id or self.default_library(workspace_id)
        now = utc_now()
        existing = None
        if source_uri:
            existing = self.db.one(
                "SELECT * FROM documents WHERE workspace_id=? AND source_uri=? ORDER BY updated_at DESC LIMIT 1",
                (workspace_id, source_uri),
            )
        document_id = existing["id"] if existing else new_id("doc")
        with self.db.transaction() as connection:
            if existing:
                next_version = int(
                    connection.execute(
                        "SELECT COALESCE(MAX(version),0)+1 FROM document_versions WHERE document_id=?",
                        (document_id,),
                    ).fetchone()[0]
                )
                old_chunk_ids = [row[0] for row in connection.execute("SELECT id FROM chunks WHERE document_id=?", (document_id,))]
                if old_chunk_ids:
                    connection.executemany("DELETE FROM chunks_fts WHERE chunk_id=?", [(item,) for item in old_chunk_ids])
                connection.execute("DELETE FROM chunks WHERE document_id=?", (document_id,))
                connection.execute(
                    """UPDATE documents SET title=?,source_type=?,content_hash=?,mime_type=?,status='READY',
                    is_gold=?,metadata_json=?,updated_at=? WHERE id=?""",
                    (title, source_type, digest, mime_type, int(is_gold), json_dumps(metadata or {}), now, document_id),
                )
            else:
                next_version = 1
                connection.execute(
                    """INSERT INTO documents
                    (id,workspace_id,library_id,title,source_type,source_uri,content_hash,mime_type,status,is_gold,metadata_json,created_at,updated_at)
                    VALUES(?,?,?,?,?,?,?,?,'READY',?,?,?,?)""",
                    (
                        document_id,
                        workspace_id,
                        library,
                        title[:200],
                        source_type,
                        source_uri,
                        digest,
                        mime_type,
                        int(is_gold),
                        json_dumps(metadata or {}),
                        now,
                        now,
                    ),
                )
            connection.execute(
                "INSERT INTO document_versions(id,document_id,version,content,content_hash,created_at) VALUES(?,?,?,?,?,?)",
                (new_id("docv"), document_id, next_version, clean, digest, now),
            )
            for ordinal, piece in enumerate(chunk_text(clean)):
                chunk_id = new_id("chk")
                connection.execute(
                    """INSERT INTO chunks(id,document_id,workspace_id,ordinal,content,token_estimate,created_at)
                    VALUES(?,?,?,?,?,?,?)""",
                    (chunk_id, document_id, workspace_id, ordinal, piece, max(1, len(piece) // 2), now),
                )
                connection.execute(
                    "INSERT INTO chunks_fts(chunk_id,workspace_id,content) VALUES(?,?,?)",
                    (chunk_id, workspace_id, piece),
                )
        return self.db.one("SELECT * FROM documents WHERE id=?", (document_id,)) or {}

    def import_bytes(self, workspace_id: str, filename: str, data: bytes, library_id: str | None = None, is_gold: bool = False) -> list[dict[str, Any]]:
        if Path(filename).suffix.lower() == ".zip":
            results: list[dict[str, Any]] = []
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                for info in archive.infolist():
                    item = Path(info.filename)
                    if info.is_dir() or item.is_absolute() or ".." in item.parts or item.suffix.lower() not in SUPPORTED_SUFFIXES:
                        continue
                    if info.file_size > 20 * 1024 * 1024:
                        continue
                    results.extend(self.import_bytes(workspace_id, item.name, archive.read(info), library_id, is_gold))
            if not results:
                raise ValueError("压缩包中没有可导入的受支持文件")
            return results
        text, mime = extract_text(filename, data)
        saved_name = f"{new_id('file')}_{safe_filename(filename)}"
        (UPLOAD_DIR / saved_name).write_bytes(data)
        return [
            self.import_text(
                workspace_id,
                Path(filename).stem,
                text,
                "file",
                saved_name,
                library_id,
                mime,
                is_gold,
                {"original_filename": safe_filename(filename)},
            )
        ]

    def import_url(self, workspace_id: str, url: str, library_id: str | None = None, is_gold: bool = False) -> dict[str, Any]:
        final_url, title, text = fetch_public_page(url)
        return self.import_text(workspace_id, title, text, "webpage", final_url, library_id, "text/html", is_gold)

    def import_folder(self, workspace_id: str, path: str, library_id: str | None = None) -> list[dict[str, Any]]:
        folder = Path(path).expanduser().resolve()
        if not folder.is_dir():
            raise ValueError("文件夹不存在")
        results: list[dict[str, Any]] = []
        for item in folder.rglob("*"):
            if item.is_file() and item.suffix.lower() in SUPPORTED_SUFFIXES and item.stat().st_size <= 20 * 1024 * 1024:
                results.extend(self.import_bytes(workspace_id, item.name, item.read_bytes(), library_id))
        return results

    def retrieve(self, workspace_id: str, query: str, limit: int = 8, gold_only: bool = False) -> list[dict[str, Any]]:
        segments = [term for term in re.split(r"[\s,，。！？:：;；、的了是在有和与]+", query) if len(term) >= 2]
        terms: list[str] = []
        for segment in segments:
            terms.append(segment)
            if re.fullmatch(r"[\u3400-\u9fff]+", segment) and len(segment) > 4:
                terms.extend(segment[index : index + 4] for index in range(len(segment) - 3))
        terms = list(dict.fromkeys(terms))[:24]
        candidates: dict[str, dict[str, Any]] = {}
        safe_terms = [term.replace('"', "") for term in terms]
        fts_query = " OR ".join(f'"{term}"' for term in safe_terms)
        if fts_query:
            try:
                rows = self.db.all(
                    """SELECT c.id,c.content,c.document_id,d.title,d.source_uri,d.is_gold,bm25(chunks_fts) AS rank
                    FROM chunks_fts JOIN chunks c ON c.id=chunks_fts.chunk_id
                    JOIN documents d ON d.id=c.document_id
                    WHERE chunks_fts MATCH ? AND c.workspace_id=?
                    ORDER BY rank LIMIT ?""",
                    (fts_query, workspace_id, max(limit * 3, 20)),
                )
                for row in rows:
                    if gold_only and not row["is_gold"]:
                        continue
                    row["score"] = 1 / (1 + abs(float(row.get("rank") or 0)))
                    candidates[row["id"]] = row
            except Exception:
                pass
        like_clauses = " OR ".join("c.content LIKE ?" for _ in terms) or "c.content LIKE ?"
        like_params = [f"%{term}%" for term in terms] or [f"%{query}%"]
        rows = self.db.all(
            f"""SELECT c.id,c.content,c.document_id,d.title,d.source_uri,d.is_gold
            FROM chunks c JOIN documents d ON d.id=c.document_id
            WHERE c.workspace_id=? AND ({like_clauses})
            ORDER BY d.is_gold DESC,d.updated_at DESC LIMIT ?""",
            (workspace_id, *like_params, max(limit * 3, 20)),
        )
        lowered_terms = [term.lower() for term in terms]
        for row in rows:
            if gold_only and not row["is_gold"]:
                continue
            haystack = row["content"].lower()
            lexical = sum(haystack.count(term) for term in lowered_terms)
            row["score"] = max(float(row.get("score") or 0), min(1.0, lexical / max(1, len(lowered_terms) * 2)))
            candidates[row["id"]] = row
        ranked = sorted(candidates.values(), key=lambda row: (float(row.get("score") or 0), int(row.get("is_gold") or 0)), reverse=True)
        return ranked[:limit]

    def build_embeddings(self, workspace_id: str, limit: int = 100) -> dict[str, int]:
        rows = self.db.all(
            "SELECT id,content FROM chunks WHERE workspace_id=? AND embedding_json IS NULL LIMIT ?",
            (workspace_id, limit),
        )
        if not rows:
            return {"updated": 0}
        vectors = self.models.embeddings(workspace_id, [row["content"] for row in rows])
        self.db.executemany(
            "UPDATE chunks SET embedding_json=? WHERE id=?",
            [(json_dumps(vector), row["id"]) for row, vector in zip(rows, vectors)],
        )
        return {"updated": len(vectors)}


knowledge_service = KnowledgeService()
